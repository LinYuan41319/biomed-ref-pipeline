import tempfile
import unittest
from pathlib import Path

from docx import Document

from biomed_ref_pipeline.docx_tools import validate_final_docx, write_review_docx
from biomed_ref_pipeline.planning import build_citation_plan, parse_pmids_cell
from biomed_ref_pipeline.scoring import classify_evidence, score_references


class PlanningScoringDocxTests(unittest.TestCase):
    def test_parse_pmids_cell(self):
        self.assertEqual(parse_pmids_cell("12345678; 23456789, 12345678"), ["12345678", "23456789"])

    def test_build_citation_plan(self):
        groups = [{"source": "x", "location": "p:1", "raw_text": "(PMID: 12345678 23456789)", "pmids": "12345678 23456789"}]
        metadata = [{"pmid": "12345678", "title": "A"}, {"pmid": "23456789", "title": "B"}]
        zotero = [{"pmid": "12345678", "in_zotero": "yes"}, {"pmid": "23456789", "in_zotero": "no"}]
        scores = [{"pmid": "12345678", "evidence_type": "experimental_mechanism", "priority_score": "70"}]
        plan = build_citation_plan(groups, metadata_rows=metadata, zotero_rows=zotero, scored_rows=scores)
        self.assertEqual(plan[0]["citation_id"], "CITE-0001")
        self.assertEqual(plan[0]["recommended_action"], "import_missing_to_zotero")
        self.assertIn("multi-source", plan[0]["notes"])

    def test_score_references(self):
        rows = [
            {
                "pmid": "1",
                "title": "Single-cell transcriptomic analysis of heart failure",
                "abstract": "single-cell transcriptomic association",
                "publication_types": "Journal Article",
                "year": "2025",
            }
        ]
        scored = score_references(rows, [{"pmid": "1", "jcr_quartile": "Q1", "jcr_5yr_if": "10"}], current_year=2026)
        self.assertEqual(scored[0]["evidence_type"], "omics_or_bioinformatics")
        self.assertEqual(scored[0]["priority_tier"], "high")

    def test_classify_review(self):
        evidence, note = classify_evidence({"publication_types": "Review", "title": "Mechanisms", "abstract": ""})
        self.assertEqual(evidence, "review_or_meta_analysis")
        self.assertIn("primary mechanism", note)

    def test_review_and_validate_docx(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "source.docx"
            review = Path(tmp) / "review.docx"
            doc = Document()
            doc.add_paragraph("Claim (PMID: 12345678).")
            doc.save(source)

            plan = [
                {
                    "citation_id": "CITE-0001",
                    "location": "paragraph:1",
                    "pmids": "12345678",
                    "pubmed_status": "verified",
                    "zotero_status": "available",
                    "recommended_action": "ready_for_zotero_insert",
                }
            ]
            write_review_docx(source, plan, review)
            self.assertTrue(review.exists())

            report = validate_final_docx(source, ["12345678"])
            self.assertIn("Remaining visible PMID groups: 1", report)
            self.assertIn("Visible PMID placeholders remain", report)


if __name__ == "__main__":
    unittest.main()
